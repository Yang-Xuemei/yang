"""Data management router: CRUD for the 8 core tables plus upload/download."""
from __future__ import annotations

import io
import json
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File
from fastapi.responses import StreamingResponse

from app.auth import get_current_user, require_admin, write_audit_log
from app.database import get_connection
from app.schemas import (
    DataStatusResponse,
    DemandItem,
    HoldingCostItem,
    InventoryItem,
    ShortageCostItem,
    SKUItem,
    TransportCostItem,
    WarehouseCapacityItem,
    WarehouseItem,
)

router = APIRouter(prefix="/api/data", tags=["data"])


# ============================
# Generic helpers
# ============================
TABLE_SPECS = {
    "skus": {
        "columns": ["sku_code", "sku_name", "sku_desc"],
        "key": "sku_code",
    },
    "warehouses": {
        "columns": ["warehouse_code", "warehouse_name", "warehouse_type", "address", "city", "longitude", "latitude", "status"],
        "key": "warehouse_code",
    },
    "inventory": {
        "columns": ["sku_code", "warehouse_code", "quantity", "last_counted_at"],
        "key": ("sku_code", "warehouse_code"),
    },
    "demand_forecast": {
        "columns": ["sku_code", "warehouse_code", "expected_demand", "accuracy"],
        "key": ("sku_code", "warehouse_code"),
    },
    "transport_cost": {
        "columns": ["source_warehouse", "target_warehouse", "unit_cost"],
        "key": ("source_warehouse", "target_warehouse"),
    },
    "holding_cost": {
        "columns": ["sku_code", "warehouse_code", "unit_cost"],
        "key": ("sku_code", "warehouse_code"),
    },
    "shortage_cost": {
        "columns": ["sku_code", "warehouse_code", "unit_cost"],
        "key": ("sku_code", "warehouse_code"),
    },
    "warehouse_capacity": {
        "columns": ["warehouse_code", "max_capacity", "current_usage", "usage_rate"],
        "key": "warehouse_code",
    },
}


def _row_to_dict(row, columns):
    d = {}
    for c in columns:
        v = row[c] if c in row.keys() else None
        d[c] = v
    return d


# ============================
# Per-table CRUD
# ============================
def _build_table_routes(table_name: str, item_model, spec):
    sub = APIRouter()

    @sub.get("")
    def list_rows(user: dict = Depends(get_current_user)):
        conn = get_connection()
        rows = conn.execute(f"SELECT * FROM {table_name} ORDER BY rowid").fetchall()
        return [_row_to_dict(r, spec["columns"]) for r in rows]

    @sub.post("")
    async def create_row(request: Request, user: dict = Depends(require_admin)):
        payload = await request.json()
        conn = get_connection()
        # Keep only valid columns
        data = {k: payload.get(k) for k in spec["columns"] if k in payload}
        if not data:
            raise HTTPException(status_code=400, detail="empty_payload")
        placeholders = ", ".join(["?"] * len(data))
        cols = ", ".join(data.keys())
        try:
            conn.execute(
                f"INSERT INTO {table_name} ({cols}) VALUES ({placeholders})",
                list(data.values()),
            )
            conn.commit()
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))
        write_audit_log(user["id"], "create", table_name, detail=json.dumps(data, ensure_ascii=False, default=str))
        return {"status": "created"}

    @sub.put("/{key_value}")
    async def update_row(key_value: str, request: Request, user: dict = Depends(require_admin)):
        payload = await request.json()
        conn = get_connection()
        data = {k: payload.get(k) for k in spec["columns"] if k in payload}
        if not data:
            raise HTTPException(status_code=400, detail="empty_payload")
        if isinstance(spec["key"], tuple):
            parts = key_value.split("__", 1)
            if len(parts) != 2:
                raise HTTPException(status_code=400, detail="invalid_key")
            where = f"{spec['key'][0]} = ? AND {spec['key'][1]} = ?"
            vals = list(data.values()) + list(parts)
        else:
            where = f"{spec['key']} = ?"
            vals = list(data.values()) + [key_value]
        set_clause = ", ".join(f"{k} = ?" for k in data.keys())
        conn.execute(
            f"UPDATE {table_name} SET {set_clause} WHERE {where}",
            vals,
        )
        conn.commit()
        write_audit_log(user["id"], "update", table_name, target_id=key_value)
        return {"status": "updated"}

    @sub.delete("/{key_value}")
    def delete_row(key_value: str, user: dict = Depends(require_admin)):
        conn = get_connection()
        if isinstance(spec["key"], tuple):
            parts = key_value.split("__", 1)
            where = f"{spec['key'][0]} = ? AND {spec['key'][1]} = ?"
            conn.execute(f"DELETE FROM {table_name} WHERE {where}", parts)
        else:
            conn.execute(f"DELETE FROM {table_name} WHERE {spec['key']} = ?", (key_value,))
        conn.commit()
        write_audit_log(user["id"], "delete", table_name, target_id=key_value)
        return {"status": "deleted"}

    return sub


# SKU
router.include_router(_build_table_routes("skus", SKUItem, TABLE_SPECS["skus"]), prefix="/skus")
router.include_router(_build_table_routes("warehouses", WarehouseItem, TABLE_SPECS["warehouses"]), prefix="/warehouses")
router.include_router(_build_table_routes("inventory", InventoryItem, TABLE_SPECS["inventory"]), prefix="/inventory")
router.include_router(_build_table_routes("demand_forecast", DemandItem, TABLE_SPECS["demand_forecast"]), prefix="/demand")
router.include_router(_build_table_routes("transport_cost", TransportCostItem, TABLE_SPECS["transport_cost"]), prefix="/transport_cost")
router.include_router(_build_table_routes("holding_cost", HoldingCostItem, TABLE_SPECS["holding_cost"]), prefix="/holding_cost")
router.include_router(_build_table_routes("shortage_cost", ShortageCostItem, TABLE_SPECS["shortage_cost"]), prefix="/shortage_cost")
router.include_router(_build_table_routes("warehouse_capacity", WarehouseCapacityItem, TABLE_SPECS["warehouse_capacity"]), prefix="/warehouse_capacity")


# ============================
# Data status
# ============================
@router.get("/status", response_model=List[DataStatusResponse])
def data_status(user: dict = Depends(get_current_user)):
    conn = get_connection()
    results = []
    for tbl, spec in TABLE_SPECS.items():
        row = conn.execute(f"SELECT COUNT(*) as c, MAX(updated_at) as u FROM {tbl}").fetchone()
        count = row["c"] if row else 0
        updated_at = row["u"] if row else None
        if count == 0:
            completeness = "❌严重缺失"
        elif tbl in ("skus", "warehouses"):
            completeness = "✅完整" if count >= 3 else "⚠️部分缺失"
        else:
            completeness = "✅完整" if count >= 6 else ("⚠️部分缺失" if count >= 3 else "❌严重缺失")
        results.append(DataStatusResponse(
            table_name=tbl,
            record_count=count,
            completeness=completeness,
            updated_at=updated_at,
        ))
    return results


# ============================
# Excel upload / download
# ============================
@router.get("/template")
def download_template(user: dict = Depends(get_current_user)):
    """Generate Excel template for download."""
    from openpyxl import Workbook
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    for tbl, spec in TABLE_SPECS.items():
        ws = wb.create_sheet(title=tbl)
        for idx, col in enumerate(spec["columns"]):
            ws.cell(row=1, column=idx + 1, value=col)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template.xlsx"},
    )


@router.get("/export")
def export_data(user: dict = Depends(get_current_user)):
    """Export all data as Excel."""
    from openpyxl import Workbook
    conn = get_connection()
    wb = Workbook()
    wb.remove(wb.active)
    for tbl, spec in TABLE_SPECS.items():
        ws = wb.create_sheet(title=tbl)
        for idx, col in enumerate(spec["columns"]):
            ws.cell(row=1, column=idx + 1, value=col)
        rows = conn.execute(f"SELECT * FROM {tbl}").fetchall()
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, col in enumerate(spec["columns"]):
                ws.cell(row=r_idx, column=c_idx + 1, value=row[col] if col in row.keys() else None)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=inventory_data.xlsx"},
    )


@router.post("/upload")
async def upload_data(
    file: UploadFile = File(...),
    user: dict = Depends(require_admin),
):
    """Upload Excel file to replace data. Rolls back entirely on any error."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="only_xlsx_supported")
    contents = await file.read()
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="file_too_large")

    import pandas as pd
    try:
        xls = pd.ExcelFile(io.BytesIO(contents), engine="openpyxl")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"invalid_excel: {str(e)[:100]}")

    errors = []
    conn = get_connection()

    # Validate all sheets first, then apply in a single transaction
    parsed = {}
    for sheet_name in xls.sheet_names:
        if sheet_name not in TABLE_SPECS:
            continue
        spec = TABLE_SPECS[sheet_name]
        try:
            df = pd.read_excel(xls, sheet_name=sheet_name)
        except Exception as e:
            errors.append(f"[{sheet_name}] 读取失败: {str(e)[:100]}")
            continue
        # Check required columns
        missing = [c for c in spec["columns"] if c not in df.columns]
        if missing:
            errors.append(f"[{sheet_name}] 缺少列: {missing}")
            continue
        # Keep only declared columns and filter empty rows
        df = df[spec["columns"]].dropna(how="all")
        parsed[sheet_name] = df

    if errors:
        raise HTTPException(status_code=400, detail={"errors": errors})

    # Apply: delete existing rows and insert new data
    try:
        for tbl, df in parsed.items():
            conn.execute(f"DELETE FROM {tbl}")
            for _, row in df.iterrows():
                values = [row[c] if pd.notna(row[c]) else None for c in TABLE_SPECS[tbl]["columns"]]
                placeholders = ", ".join(["?"] * len(values))
                cols = ", ".join(TABLE_SPECS[tbl]["columns"])
                conn.execute(f"INSERT INTO {tbl} ({cols}) VALUES ({placeholders})", values)
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=f"upload_failed: {str(e)[:100]}")

    write_audit_log(user["id"], "upload", "data", detail=f"sheets={list(parsed.keys())}")
    return {"status": "uploaded", "sheets": list(parsed.keys())}

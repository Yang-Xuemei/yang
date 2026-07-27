"""Compare multiple solve history results."""
from __future__ import annotations

import json
from typing import List

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from app.auth import get_current_user
from app.database import get_connection
from app.routers.history import _to_history_dict
from app.schemas import CompareRequest

router = APIRouter(prefix="/api/compare", tags=["compare"])


@router.post("", response_model=dict)
def compare(payload: CompareRequest, user: dict = Depends(get_current_user)):
    conn = get_connection()
    records = []
    for hid in payload.history_ids:
        row = conn.execute("SELECT * FROM solve_history WHERE id = ?", (hid,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail=f"history_{hid}_not_found")
        records.append(_to_history_dict(row))

    # Compute summary
    if records:
        best_cost_idx = min(range(len(records)), key=lambda i: records[i]["total_cost"])
        best_sl_idx = max(range(len(records)), key=lambda i: records[i]["service_level"])
        baseline_cost = next((r["total_cost"] for r in records if r["algorithm_type"] == "baseline"), None)
        improvements = []
        for r in records:
            if baseline_cost and baseline_cost > 0:
                imp = (baseline_cost - r["total_cost"]) / baseline_cost
                improvements.append(round(imp * 100, 2))
            else:
                improvements.append(None)
        summary = {
            "best_cost": {"index": best_cost_idx, "value": records[best_cost_idx]["total_cost"]},
            "best_service_level": {"index": best_sl_idx, "value": records[best_sl_idx]["service_level"]},
            "improvements_vs_baseline": improvements,
        }
    else:
        summary = {}

    return {"records": records, "summary": summary}


@router.post("/export")
def export_compare(payload: CompareRequest, user: dict = Depends(get_current_user)):
    """Export comparison as Excel."""
    conn = get_connection()
    records = []
    for hid in payload.history_ids:
        row = conn.execute("SELECT * FROM solve_history WHERE id = ?", (hid,)).fetchone()
        if row is None:
            continue
        records.append(_to_history_dict(row))

    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "对比结果"
    headers = ["求解编号", "算法类型", "参数场景", "总成本", "运输成本", "缺货成本", "持仓成本", "服务水平", "调拨次数", "调拨总量", "求解时间"]
    for idx, h in enumerate(headers, 1):
        ws.cell(row=1, column=idx, value=h)

    for r_idx, r in enumerate(records, start=2):
        res = r.get("result", {}) or {}
        ws.cell(row=r_idx, column=1, value=r["solve_name"])
        ws.cell(row=r_idx, column=2, value=r["algorithm_type"])
        ws.cell(row=r_idx, column=3, value=r.get("scenario_name") or "")
        ws.cell(row=r_idx, column=4, value=r["total_cost"])
        ws.cell(row=r_idx, column=5, value=res.get("transport_cost", 0))
        ws.cell(row=r_idx, column=6, value=res.get("shortage_cost", 0))
        ws.cell(row=r_idx, column=7, value=res.get("holding_cost", 0))
        ws.cell(row=r_idx, column=8, value=r["service_level"])
        ws.cell(row=r_idx, column=9, value=res.get("transfer_count", 0))
        ws.cell(row=r_idx, column=10, value=res.get("total_transfer_qty", 0))
        ws.cell(row=r_idx, column=11, value=r["created_at"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=compare.xlsx"},
    )
